#!/usr/bin/env python3
"""
Создаёт DOCX-шаблон «Рабочая карта обследования EXPERTUS» с плейсхолдерами {{...}}.
Оформление: Times New Roman, чёткие блоки (реквизиты, условия, оборудование, документация, дефектная ведомость, эскиз).
Запуск из корня backend: python scripts/generate_inspection_template.py
"""
import os
import sys

# Add app to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "app", "templates")
OUTPUT_PATH = os.path.join(TEMPLATE_DIR, "inspection_card_template.docx")

FONT_NAME = "Times New Roman"
FONT_SIZE_BODY = Pt(11)
FONT_SIZE_SMALL = Pt(10)
FONT_SIZE_TITLE = Pt(14)
FONT_SIZE_SUBTITLE = Pt(12)
FONT_SIZE_FOOTER = Pt(8)


def _section_heading(doc: Document, text: str) -> None:
    """Заголовок раздела: жирный, 11 pt, отступ сверху."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run(text)
    run.bold = True
    run.font.size = FONT_SIZE_BODY
    run.font.name = FONT_NAME


def main() -> None:
    os.makedirs(TEMPLATE_DIR, exist_ok=True)
    doc = Document()
    style = doc.styles["Normal"]
    if style.font:
        style.font.name = FONT_NAME
        style.font.size = FONT_SIZE_BODY

    # ---------- Шапка ----------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Приложение 3").font.size = FONT_SIZE_SMALL
    p.add_run("\n").font.name = FONT_NAME
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ТОО «EXPERTUS»").font.size = FONT_SIZE_SUBTITLE
    p.add_run("\n").font.name = FONT_NAME
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Ф-EXP-ДП-ИЛ-7.8-03").font.size = FONT_SIZE_SMALL
    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Рабочая карта обследования").font.size = FONT_SIZE_TITLE
    p.runs[0].font.name = FONT_NAME
    doc.add_paragraph()
    doc.add_paragraph()

    # ---------- Реквизиты объекта ----------
    _section_heading(doc, "1. Реквизиты и условия")

    doc.add_paragraph("Заказчик (Договор/проект): {{customer}}")
    doc.add_paragraph("Наименование объекта обследования: {{object_name}}")
    doc.add_paragraph("Место нахождения: {{location}}")
    doc.add_paragraph("Инв. № {{inventory_number}}   Зав. № {{factory_number}}")
    doc.add_paragraph("Ф.И.О. исполнителя: {{inspector_name}}   Дата: {{inspection_date}}")
    doc.add_paragraph()
    doc.add_paragraph(
        "Температура воздуха, °C: {{temperature_c}}   "
        "Влажность воздуха, %: {{humidity_percent}}   "
        "Уровень освещенности, лк: {{illumination_lux}}"
    )
    doc.add_paragraph()

    # ---------- Оборудование (таблица) ----------
    _section_heading(doc, "2. Применяемое оборудование")

    table = doc.add_table(rows=6, cols=2)
    table.autofit = False
    for col_idx in range(2):
        for row in table.rows:
            row.cells[col_idx].width = Cm(7.5)
    # Заголовки
    h0, h1 = table.rows[0].cells
    h0.text = "Наименование"
    h1.text = "Зав. №"
    for c in (h0, h1):
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = FONT_SIZE_BODY
                r.font.name = FONT_NAME
    # Строки с плейсхолдерами
    rows_text = [
        ("Комплект ВИК", "{{equipment_vik}}"),
        ("Комплект пенетрантов ПВК", "{{equipment_penetrant}}"),
        ("Ультразвуковой дефектоскоп", "{{equipment_defectoscope}}"),
        ("Ультразвуковой толщиномер", "{{equipment_thickness_gauge}}"),
        ("Твердомер (ультразвуковой/динамический)", "{{equipment_hardness}}"),
    ]
    for i, (name, ph) in enumerate(rows_text, start=1):
        table.rows[i].cells[0].text = name
        table.rows[i].cells[1].text = ph
    doc.add_paragraph()

    # ---------- Состояние документации ----------
    _section_heading(doc, "3. Состояние рабочей документации на объект испытания/контроля")

    doc.add_paragraph("{{documentation_line}}")
    doc.add_paragraph()

    # ---------- Дефектная ведомость ----------
    _section_heading(doc, "4. Дефектная ведомость / примечания")

    for _ in range(4):
        doc.add_paragraph("_" * 100)
    doc.add_paragraph("_" * 50)
    doc.add_paragraph("{{defects_notes}}")
    doc.add_paragraph()

    # ---------- Эскиз ----------
    _section_heading(doc, "5. Эскиз, схема объекта испытаний")

    doc.add_paragraph("{{sketch_note}}")
    doc.add_paragraph()
    doc.add_paragraph()

    # ---------- Подвал ----------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run(
        "© Настоящий документ не может быть полностью или частично воспроизведен, "
        "тиражирован и распространен в качестве официального издания без разрешения руководства ТОО «Expertus»"
    )
    run.font.size = FONT_SIZE_FOOTER
    run.font.name = FONT_NAME

    doc.save(OUTPUT_PATH)
    print(f"Template saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
